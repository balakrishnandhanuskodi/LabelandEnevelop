#!/usr/bin/env python3
"""
Entrainment 26 - Label & Envelope PDF generator.

Reads an Excel/CSV guest list and writes two print-ready PDFs:
  * <name>_labels.pdf    - 12 labels per A4 page (2 x 6), sizes from config.json
  * <name>_envelopes.pdf - one envelope per guest (DL 220 x 110 mm); prints only in the
                           top-right print area set in config.json

Usage:
  python generate.py data.xlsx                 # writes to ./output
  python generate.py data.xlsx --out prints    # custom output folder
  python generate.py data.xlsx --guides        # draw label outlines (for test prints)
  python generate.py data.xlsx --config my.json

Print at "Actual size" / 100% scale, no "fit to page".
"""
import argparse
import csv
import hashlib
import json
import re
import sys
import urllib.parse
import urllib.request
from pathlib import Path

import qrcode
from openpyxl import load_workbook
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = Path(__file__).resolve().parent
FONT, FONT_BOLD = "Helvetica", "Helvetica-Bold"


# ---------------------------------------------------------------- fonts
def register_fonts():
    """Caladea is metric-compatible with Cambria (the font the HTML portal used)."""
    global FONT, FONT_BOLD
    reg, bold = HERE / "fonts" / "Caladea-Regular.ttf", HERE / "fonts" / "Caladea-Bold.ttf"
    if reg.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont("Caladea", str(reg)))
        pdfmetrics.registerFont(TTFont("Caladea-Bold", str(bold)))
        FONT, FONT_BOLD = "Caladea", "Caladea-Bold"


# ---------------------------------------------------------------- data
def clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_rows(path: Path):
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
    else:
        # not read-only, so Excel hyperlink targets are available
        ws = load_workbook(path, data_only=True).worksheets[0]
        rows = [[(c.value, c.hyperlink.target if c.hyperlink and c.hyperlink.target else "") for c in r]
                for r in ws.iter_rows()]
    if path.suffix.lower() == ".csv":
        rows = [[(v, "") for v in r] for r in rows]
    rows = [r for r in rows if any(clean(v) or l for v, l in r)]
    if not rows:
        sys.exit(f"No data found in {path}")
    headers = [clean(v) for v, _ in rows[0]]
    records = []
    for r in rows[1:]:
        rec = {h: clean(r[i][0]) if i < len(r) else "" for i, h in enumerate(headers) if h}
        rec["__links"] = {h: str(r[i][1]).strip() for i, h in enumerate(headers) if h and i < len(r) and r[i][1]}
        records.append(rec)
    return headers, records


def map_columns(headers, candidates):
    """Same idea as the portal's auto-detect: exact match first, then 'contains'."""
    mapping = {}
    lower = {h.lower(): h for h in headers if h}
    for field, names in candidates.items():
        found = next((lower[n.lower()] for n in names if n.lower() in lower), None)
        if not found:
            found = next((h for n in names for h in headers
                          if h and n.lower() in h.lower()), None)
        mapping[field] = found
    return mapping


def fill_pattern(pattern, rec):
    return re.sub(r"\{([^}]+)\}", lambda m: str(rec.get(m.group(1), "")), pattern).strip()


def guest(rec, cols, qr_pattern, qr_mode="text"):
    g = {k: (rec.get(c, "") if c else "") for k, c in cols.items()}
    # QR = the QR column's value exactly as in the sheet (hyperlink address if the cell is a link).
    # Nothing is added; an empty cell gives no QR. The pattern is only used when there is no QR column.
    if cols.get("qr"):
        g["qr"] = rec["__links"].get(cols["qr"]) or rec.get(cols["qr"], "")
    else:
        g["qr"] = fill_pattern(qr_pattern, rec) if qr_pattern else ""
    g["qr_mode"] = qr_mode if cols.get("qr") else "text"
    g["resort_line"] = " - ".join(x for x in (g["resort"], g["building"]) if x)
    parts = []
    if g["floor"]:
        parts.append(g['floor'])
    if g["room"]:
        parts.append(f"Room # {g['room']}")
    g["floor_room"] = " - ".join(parts)
    return g


# ---------------------------------------------------------------- original QR images
IMG_RE = re.compile(r"\.(jpe?g|png|gif|webp)(?:[?#]|$)", re.I)
QR_IMAGES = {}      # url -> local file path (embedded exactly as is)


def is_image_link(v):
    return bool(re.match(r"https?://", v, re.I) and IMG_RE.search(v))


def url_file_name(url):
    return urllib.parse.unquote(urllib.parse.urlparse(url).path.rsplit("/", 1)[-1]).lower()


def load_qr_images(urls, cache_dir: Path, local_dir: Path | None):
    """Find each QR image: a matching file in --qr-images, else download (cached)."""
    local = {}
    if local_dir:
        local = {p.name.lower(): p for p in local_dir.iterdir() if p.is_file()}
    cache_dir.mkdir(parents=True, exist_ok=True)
    missing = []
    for url in urls:
        name = url_file_name(url)
        if name in local:
            QR_IMAGES[url] = local[name]
            continue
        ext = Path(name).suffix or ".img"
        dest = cache_dir / (hashlib.sha1(url.encode()).hexdigest()[:16] + ext)
        if not dest.exists():
            try:
                req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
                with urllib.request.urlopen(req, timeout=30) as r:
                    dest.write_bytes(r.read())
            except Exception as e:
                missing.append((url, str(e)))
                continue
        QR_IMAGES[url] = dest
    return missing


# ---------------------------------------------------------------- drawing helpers
_qr_cache = {}


def qr_image(text):
    if text not in _qr_cache:
        q = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, border=0, box_size=10)
        q.add_data(text)
        q.make(fit=True)
        _qr_cache[text] = ImageReader(q.make_image(fill_color="black", back_color="white").get_image())
    return _qr_cache[text]


def fit_size(text, font, size, max_w, min_size=7):
    """Shrink font until text fits max_w (points)."""
    while size > min_size and pdfmetrics.stringWidth(text, font, size) > max_w:
        size -= 0.5
    return size


def wrap_two_lines(text, font, size, max_w, min_size):
    """Return (lines, size): one line if it fits, else split into two balanced lines."""
    if pdfmetrics.stringWidth(text, font, size) <= max_w or " " not in text:
        return [text], fit_size(text, font, size, max_w, min_size)
    words = text.split()
    best = min(range(1, len(words)),
               key=lambda i: max(pdfmetrics.stringWidth(" ".join(words[:i]), font, size),
                                 pdfmetrics.stringWidth(" ".join(words[i:]), font, size)))
    lines = [" ".join(words[:best]), " ".join(words[best:])]
    s = min(fit_size(l, font, size, max_w, min_size) for l in lines)
    return lines, s


def draw_qr_image(c, url, x, y, size):
    """Place the original QR image unchanged (fitted inside the square if not square)."""
    path = QR_IMAGES.get(url)
    if not path:
        c.saveState()
        c.setStrokeColorRGB(0.6, 0.6, 0.6); c.setDash(2, 2); c.setLineWidth(0.5)
        c.rect(x, y, size, size)
        c.setFont(FONT, 5); c.setFillColorRGB(0.4, 0.4, 0.4)
        c.drawCentredString(x + size / 2, y + size / 2, "QR image missing")
        c.restoreState()
        return
    img = ImageReader(str(path))
    w, h = img.getSize()
    k = min(size / w, size / h)
    c.drawImage(str(path), x + (size - w * k) / 2, y + (size - h * k) / 2, w * k, h * k)


def draw_stack(c, items, cx, top, bottom, gap=0, start=None):
    """items: list of (kind, payload, height_pt, font, size).
    Centered vertically between top/bottom (or starting at y=start), with `gap` between items.
    Returns the y of the stack's top edge."""
    total = sum(i[2] for i in items) + gap * max(len(items) - 1, 0)
    y = start if start is not None else bottom + (top - bottom + total) / 2
    top_edge = y
    for n, (kind, payload, h, font, size) in enumerate(items):
        if n:
            y -= gap
        if kind == "text":
            c.setFont(font, size)
            c.drawCentredString(cx, y - h + (h - size * 0.7) / 2, payload)
        elif kind == "qr":
            c.drawImage(qr_image(payload), cx - h / 2, y - h, h, h)
        elif kind == "qrimg":
            draw_qr_image(c, payload, cx - h / 2, y - h, h)
        y -= h
    return top_edge


def text_item(text, font, size, max_w, gap=1.15):
    size = fit_size(text, font, size, max_w)
    return ("text", text, size * gap, font, size)


# ---------------------------------------------------------------- labels
def printed_title(cfg):
    return cfg["event_title"] if cfg.get("print_event_title") else ""


def draw_label(c, g, x, y, w, h, L, title, guides):
    fs = L["font_sizes_pt"]
    pad = L["padding_mm"] * mm
    if guides:
        c.setStrokeColorRGB(0.75, 0.75, 0.75)
        c.setLineWidth(0.3)
        c.rect(x, y, w, h)
    ix, iy, iw, ih = x + pad, y + pad, w - 2 * pad, h - 2 * pad
    left_w = iw * 0.25
    right_x, right_w = ix + left_w + 2 * mm, iw - left_w - 2 * mm

    # left: global seq / QR / block seq, with a gap between each
    gap = L.get("stack_gap_mm", 2) * mm
    qr = min(L["qr_size_mm"] * mm, left_w, ih - (fs["global_seq"] + fs["block_seq"]) * 1.15 - 2 * gap)
    left = []
    if g["global_seq"]:
        left.append(text_item(g["global_seq"], FONT_BOLD, fs["global_seq"], left_w))
    if g["qr"]:
        left.append(("qrimg" if g.get("qr_mode") == "image" else "qr", g["qr"], qr, None, None))
    if g["block_seq"]:
        left.append(text_item(g["block_seq"], FONT_BOLD, fs["block_seq"], left_w))
    # right: (optional event title) / name / resort - building / floor - room
    right = []
    if title:
        right.append(text_item(title.upper(), FONT, fs["event"], right_w, 1.5))
    if g["name"]:
        lines, s = wrap_two_lines(g["name"], FONT_BOLD, fs["name"], right_w, 11)
        right += [("text", l, s * 1.15, FONT_BOLD, s) for l in lines]
    if g["resort_line"]:
        right.append(text_item(g["resort_line"], FONT_BOLD, fs["resort"], right_w))
    if g["floor_room"]:
        right.append(text_item(g["floor_room"], FONT, fs["floor_room"], right_w))

    if L.get("anchor") == "top-right":
        # Pack the whole block into the box's top-right corner: the widest text line
        # ends at the right edge and the first line's capitals touch the top edge.
        def width(it):
            return pdfmetrics.stringWidth(it[1], it[3], it[4]) if it[0] == "text" else it[2]
        def cap_offset(items):
            return (items[0][2] - items[0][4] * 0.7) / 2 if items and items[0][0] == "text" else 0
        rw = max((width(i) for i in right), default=0)
        lw = max((width(i) for i in left), default=0)
        col_gap = L.get("column_gap_mm", 4) * mm
        top = iy + ih
        r_cx = ix + iw - rw / 2
        l_cx = ix + iw - rw - col_gap - lw / 2
        draw_stack(c, left, l_cx, top, iy, gap, start=top + cap_offset(left))
        draw_stack(c, right, r_cx, top, iy, start=top + cap_offset(left))
        return

    # Label without a QR: centre the whole content horizontally and vertically,
    # number column and text block side by side, starting level with each other.
    if not g["qr"]:
        def tw_(it):
            return pdfmetrics.stringWidth(it[1], it[3], it[4])
        lw_ = max((tw_(i) for i in left), default=0)
        rw_ = max((tw_(i) for i in right), default=0)
        cg = L.get("column_gap_mm", 4) * mm if left and right else 0
        x0 = ix + (iw - (lw_ + cg + rw_)) / 2
        block = max(sum(i[2] for i in left) + gap * max(len(left) - 1, 0), sum(i[2] for i in right))
        start = iy + (ih + block) / 2
        draw_stack(c, left, x0 + lw_ / 2, iy + ih, iy, gap, start=start)
        draw_stack(c, right, x0 + lw_ + cg + rw_ / 2, iy + ih, iy, start=start)
        return

    # default (labels): left column centred vertically, right block starts level with it
    left_top = draw_stack(c, left, ix + left_w / 2, iy + ih, iy, gap)
    draw_stack(c, right, right_x + right_w / 2, iy + ih, iy, start=left_top if left else None)


def label_copies(guests, copies):
    """Repeat each guest `copies` times in a row; only the first copy keeps the QR."""
    out = []
    for g in guests:
        out.append(g)
        for _ in range(max(1, copies) - 1):
            out.append({**g, "qr": "", "qr_slot": True})
    return out


def build_labels(guests, cfg, out, guides, copies=1):
    L = cfg["labels"]
    guests = label_copies(guests, copies)
    pw, ph = L["page_width_mm"] * mm, L["page_height_mm"] * mm
    lw, lh = L["label_width_mm"] * mm, L["label_height_mm"] * mm
    cols, rows = L["columns"], L["rows"]
    per_page = cols * rows
    c = canvas.Canvas(str(out), pagesize=(pw, ph))
    c.setTitle(f"{cfg['event_title']} - Labels")
    for i, g in enumerate(guests):
        slot = i % per_page
        if slot == 0 and i:
            c.showPage()
        r, col = divmod(slot, cols)
        x = L["margin_left_mm"] * mm + col * (lw + L["gap_x_mm"] * mm)
        y = ph - L["margin_top_mm"] * mm - (r + 1) * lh - r * L["gap_y_mm"] * mm
        draw_label(c, g, x, y, lw, lh, L, printed_title(cfg), guides)
    c.save()
    return -(-len(guests) // per_page)


# ---------------------------------------------------------------- envelopes
def deep_merge(base, over):
    out = json.loads(json.dumps(base))
    for k, v in over.items():
        out[k] = deep_merge(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else v
    return out


def envelope_types(cfg):
    """standard = cfg.envelope; every other type = standard + its overrides (e.g. leonia)."""
    types = {"standard": cfg["envelope"]}
    for name, over in cfg.get("envelope_types", {}).items():
        types[name] = deep_merge(cfg["envelope"], {k: v for k, v in over.items() if not k.startswith("_")})
    return types


def build_envelopes(guests, cfg, out, guides=False, E=None):
    """Prints only inside the envelope's print_area, using the same layout as a label."""
    E = E or cfg["envelope"]
    A = E["print_area"]
    w, h = E["width_mm"] * mm, E["height_mm"] * mm
    aw, ah = A["width_mm"] * mm, A["height_mm"] * mm
    x = (w - A["right_mm"] * mm - aw) if "right_mm" in A else A["left_mm"] * mm
    y = h - A["top_mm"] * mm - ah
    c = canvas.Canvas(str(out), pagesize=(w, h))
    c.setTitle(f"{cfg['event_title']} - Envelopes")
    for g in guests:
        draw_label(c, g, x, y, aw, ah, E, printed_title(cfg), guides)
        c.showPage()
    c.save()


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("data", type=Path, help="Excel (.xlsx) or CSV guest list")
    ap.add_argument("--config", type=Path, default=HERE / "config.json")
    ap.add_argument("--out", type=Path, default=Path("output"))
    ap.add_argument("--qr-images", type=Path, help="folder of QR image files to use instead of downloading (matched by file name)")
    ap.add_argument("--envelopes", default="all",
                    help="envelope types to build, comma-separated (standard, leonia) or 'all' (default)")
    ap.add_argument("--copies", type=int, choices=[1, 4], default=None,
                    help="labels per guest (1 or 4). With 4, only the first copy has the QR. Default: labels.copies in config.json")
    ap.add_argument("--guides", action="store_true", help="draw label / envelope print-area outlines for an alignment test print")
    a = ap.parse_args()

    cfg = json.loads(a.config.read_text(encoding="utf-8"))
    register_fonts()
    headers, records = read_rows(a.data)
    cols = map_columns(headers, cfg["columns"])
    if not cols.get("name"):
        sys.exit(f"Could not find a Name column. Headers found: {headers}")
    for f, h in cols.items():
        print(f"  {f:<11} <- {h or '(not found)'}")

    qr_mode = "text"
    if cols.get("qr"):
        vals = [(r["__links"].get(cols["qr"]) or r.get(cols["qr"], "")) for r in records]
        vals = [v for v in vals if v]
        want = cfg.get("qr_mode", "auto")
        if want == "image" or (want == "auto" and vals and sum(map(is_image_link, vals)) >= len(vals) / 2):
            qr_mode = "image"
    guests = [guest(r, cols, cfg["qr_pattern"], qr_mode) for r in records]
    guests = [g for g in guests if g["name"]]
    if qr_mode == "image":
        urls = sorted({g["qr"] for g in guests if g["qr"]})
        missing = load_qr_images(urls, HERE / ".qr_cache", a.qr_images)
        print(f"  QR images: {len(urls) - len(missing)} of {len(urls)} found (embedded exactly as is)")
        for url, err in missing[:10]:
            print(f"    missing: {url}  ({err})")
    a.out.mkdir(parents=True, exist_ok=True)
    stem = a.data.stem
    lab = a.out / f"{stem}_labels.pdf"
    copies = a.copies or cfg["labels"].get("copies", 1)
    pages = build_labels(guests, cfg, lab, a.guides, copies)
    types = envelope_types(cfg)
    wanted = list(types) if a.envelopes == "all" else [t.strip() for t in a.envelopes.split(",")]
    env_files = []
    for t in wanted:
        if t not in types:
            sys.exit(f"Unknown envelope type '{t}'. Known: {', '.join(types)}")
        env = a.out / (f"{stem}_envelopes.pdf" if t == "standard" else f"{stem}_envelopes_{t}.pdf")
        build_envelopes(guests, cfg, env, a.guides, types[t])
        env_files.append(str(env))
    print(f"{len(guests)} guests x {copies} label{'s' if copies > 1 else ''} -> {lab} ({pages} pages)")
    for f in env_files:
        print(f"  envelopes -> {f} ({len(guests)} envelopes)")


if __name__ == "__main__":
    main()
